import pptx
import json

prs = pptx.Presentation('Combined Review.pptx')
presentation_data = []

for s_idx, slide in enumerate(prs.slides):
    slide_data = {
        'slide_index': s_idx,
        'shapes': []
    }
    for sh_idx, shape in enumerate(slide.shapes):
        sh_dict = {
            'index': sh_idx,
            'id': shape.shape_id,
            'name': shape.name,
            'left': shape.left,
            'top': shape.top,
            'width': shape.width,
            'height': shape.height,
            'has_text': shape.has_text_frame,
            'has_table': shape.has_table,
        }
        if shape.has_text_frame:
            sh_dict['text'] = shape.text_frame.text
            sh_dict['paragraphs'] = []
            for p in shape.text_frame.paragraphs:
                p_dict = {
                    'text': p.text,
                    'runs': []
                }
                for r in p.runs:
                    color_hex = None
                    try:
                        if r.font.color and r.font.color.type == 1:
                            color_hex = str(r.font.color.rgb)
                    except:
                        pass
                    p_dict['runs'].append({
                        'text': r.text,
                        'font_name': r.font.name,
                        'font_size': r.font.size.pt if r.font.size else None,
                        'bold': r.font.bold,
                        'italic': r.font.italic,
                        'color': color_hex
                    })
                sh_dict['paragraphs'].append(p_dict)
        if shape.has_table:
            sh_dict['table_rows'] = []
            for row in shape.table.rows:
                sh_dict['table_rows'].append([c.text for c in row.cells])
        slide_data['shapes'].append(sh_dict)
    presentation_data.append(slide_data)

with open('full_template_details.json', 'w', encoding='utf-8') as f:
    json.dump(presentation_data, f, indent=2, ensure_ascii=False)

print("Saved full_template_details.json with", len(presentation_data), "slides")
