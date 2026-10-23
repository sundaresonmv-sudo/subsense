import pptx

prs = pptx.Presentation('Combined Review.pptx')
print("Total slides:", len(prs.slides))

for i in range(min(5, len(prs.slides))):
    slide = prs.slides[i]
    print(f"\n--- SLIDE {i+1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    color_val = None
                    try:
                        if r.font.color:
                            color_val = str(r.font.color.rgb) if r.font.color.type == 1 else str(r.font.color.theme_color)
                    except:
                        pass
                    print(f"  Shape [{shape.name}] Text: {repr(r.text[:25])} | Font: {r.font.name} | Size: {r.font.size} | Color: {color_val}")
