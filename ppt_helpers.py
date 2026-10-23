import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import copy
import os

# Colors from template
TEAL = RGBColor(0x00, 0xA8, 0x96)       # #00A896
DEEP_TEAL = RGBColor(0x02, 0x80, 0x90)  # #028090
CHARCOAL = RGBColor(0x1B, 0x2A, 0x2E)   # #1B2A2E
MUTED = RGBColor(0x7A, 0x8F, 0x8C)      # #7A8F8C
WHITE = RGBColor(0xFF, 0xFF, 0xFF)      # #FFFFFF

def set_run_style(run, font_name="Calibri", font_size=13, bold=False, italic=False, color=CHARCOAL):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def update_text_frame(tf, bullets, font_name="Calibri", font_size=13, bold=False, color=CHARCOAL, line_spacing=1.15):
    tf.clear()
    for idx, item in enumerate(bullets):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = item
        set_run_style(run, font_name=font_name, font_size=font_size, bold=bold, color=color)

def update_table_cell(cell, text, font_name="Calibri", font_size=11, bold=False, color=CHARCOAL, align=PP_ALIGN.LEFT):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    if p.runs:
        set_run_style(p.runs[0], font_name=font_name, font_size=font_size, bold=bold, color=color)

print("Helper functions ready")
